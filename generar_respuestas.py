from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(11)
style.font.color.rgb = RGBColor(0, 0, 0)

# Forzar headings a negro
for i in range(1, 4):
    h = doc.styles[f'Heading {i}']
    h.font.color.rgb = RGBColor(0, 0, 0)
    h.font.name = 'Arial'

# --- Portada ---
for _ in range(6):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('EXAMEN FINAL')
run.bold = True
run.font.size = Pt(20)
run.font.color.rgb = RGBColor(0, 0, 0)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Ingenieria de Software II')
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0, 0, 0)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Enrique Lujan')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('github.com/tunek12421/project-exam-ingenieria-software')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Rama: examen-enrique-lujan')

doc.add_page_break()

# ========================
# PARTE A
# ========================
doc.add_heading('A. Parte Teorica', level=1)

doc.add_heading('1. Monolitica vs por capas', level=2)
doc.add_paragraph(
    'La monolitica tiene todo junto en un solo bloque. La de capas separa en niveles '
    '(presentacion, negocio, datos), cada uno con su responsabilidad. Es mas facil de mantener y probar.'
)

doc.add_heading('2. Refactorizacion y funcionalidad', level=2)
doc.add_paragraph(
    'La refactorizacion mejora la estructura del codigo sin cambiar lo que hace. '
    'Si cambia el comportamiento, ya no es refactorizacion. Las pruebas deben seguir pasando igual.'
)

doc.add_heading('3. Pruebas automatizadas', level=2)
doc.add_paragraph(
    'Detectan errores rapido y sin probar a mano. Cuando cambias algo, te avisan si rompiste algo. '
    'Sirven como documentacion del comportamiento esperado.'
)

doc.add_page_break()

# ========================
# PARTE B
# ========================
doc.add_heading('B. Parte Practica', level=1)

# Act 1
doc.add_heading('Actividad 1 - Organizacion y versiones', level=2)
doc.add_paragraph(
    'Python/Flask, HTML/CSS/JS, SQLite, pytest, Docker. '
    'Rama examen-enrique-lujan con 7 commits.'
)

doc.add_paragraph('Estructura:')
p = doc.add_paragraph()
run = p.add_run(
    'backend/src/ -> config, controllers, services, models, repositories\n'
    'frontend/src/ -> components, pages, services, assets\n'
    'database/ -> schema.sql, seed.sql\n'
    'tests/ -> unitarias e integracion\n'
    'Dockerfile, docker-compose.yml, README.md'
)
run.font.name = 'Consolas'
run.font.size = Pt(9)

# Act 2
doc.add_heading('Actividad 2 - Arquitectura', level=2)
doc.add_paragraph(
    'Arquitectura por capas: Frontend -> Controllers -> Services -> Repositories -> BD. '
    'Cada capa con su responsabilidad. Modelos: TaskList y Subtask.'
)
doc.add_paragraph(
    'Mejoras posibles: migraciones con Alembic, autenticacion JWT, '
    'middleware de errores, migrar frontend a Vue/React.'
)

# Act 3
doc.add_heading('Actividad 3 - Refactorizacion', level=2)
doc.add_paragraph('Problemas encontrados y corregidos:')
smells = [
    'Conexiones a BD sin cierre seguro -> context manager (with)',
    'Query SQL repetida 3 veces -> constante _BASE_SELECT',
    'Import datetime sin usar -> eliminado',
    'Numero 200 directo en codigo -> constante TITLE_MAX_LENGTH',
    'Validacion duplicada en servicio -> metodo _find_list_or_raise()',
    'Sin validacion en frontend -> trim() + check vacio',
    'Sin aviso de exito -> notificaciones toast',
]
for s in smells:
    doc.add_paragraph(s, style='List Bullet')

# Act 4
doc.add_heading('Actividad 4 - Funcionalidad', level=2)
doc.add_paragraph('El sistema permite:')
for f in [
    'Crear tareas con titulo y descripcion',
    'Agregar subtareas dentro de cada tarea',
    'Marcar y desmarcar subtareas (toggle)',
    'Circulo de progreso segun subtareas completadas',
    'Expandir/colapsar subtareas',
    'Eliminar tareas y subtareas',
    'Validacion de campos vacios',
]:
    doc.add_paragraph(f, style='List Bullet')

# Act 5
doc.add_heading('Actividad 5 - Pruebas', level=2)
doc.add_paragraph('21 pruebas (15 unitarias + 6 integracion). Todas pasan.')
doc.add_paragraph()
doc.add_paragraph('Unitarias: validaciones de titulo, CRUD listas, CRUD subtareas, toggle on/off, calculo de progreso (0%, 25%, 100%).')
doc.add_paragraph()
doc.add_paragraph('Integracion: crear y listar, error 400 sin titulo, progreso con subtareas, toggle check/uncheck, cascade delete, error 404.')
doc.add_paragraph()
p = doc.add_paragraph()
run = p.add_run('$ python3 -m pytest -v\n21 passed in 0.39s')
run.font.name = 'Consolas'
run.font.size = Pt(9)

# Guardar
output = '/home/tunek/Descargas/DOCUMENTACION/project-exam/Examen_Final_Respuestas.docx'
doc.save(output)
print(f'Generado: {output}')
